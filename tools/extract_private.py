"""Převede list „Výsledky IB“ ze zdrojového Excelu na JSON pro db/import-private.js.

Použití (výstup patří do data/private/, který git ignoruje):
    python3 tools/extract_private.py "Srovnání univerzit 2026 – Lukas.xlsx" data/private/lukas-private.json

Rozložení listu je popsané v docs/02-analyza-excelu.md#list-výsledky-ib.
"""
import json
import re
import sys

import openpyxl

SUBJECT_CODES = {
    'Sports Science': 'SEHS',
    'Biology': 'BIO',
    'Business and Management': 'BM',
    'English': 'ENG_A',
    'Spanish': 'SPA_B',
    'Math': 'MATH_AA',
}
SCENARIO_COLS = 'CDEFG'          # kašle na to … moje predikce
ASSESSMENT_COLS = 'CDEFGHIJ'
EXAM_COLS = 'LMNO'
TERM = 'G11-S1'


def criterion(name):
    m = re.search(r'\b(KU|TL|ToL)\s*$', name.strip())
    if not m:
        return None
    return 'TL' if m.group(1) == 'ToL' else m.group(1)


def main(src, dst):
    ws = openpyxl.load_workbook(src, data_only=True)['Výsledky IB']
    cell = lambda col, row: ws[f'{col}{row}'].value

    subjects, scenarios = [], []
    for row in range(2, 8):
        code = SUBJECT_CODES[cell('B', row).strip()]
        subjects.append({'code': code, 'ismMean2025': cell('H', row), 'worldMean2025': cell('I', row), 'row': row})

    for col in SCENARIO_COLS:
        scenarios.append({
            'name': cell(col, 1).strip(),
            'corePoints': cell(col, 8),   # řádek „Extended Essay“ = body za TOK+EE
            'grades': {s['code']: cell(col, s['row']) for s in subjects},
        })

    assessments, components = [], []
    for row in range(17, ws.max_row + 1):
        name = cell('B', row)
        if not name or name.strip() not in SUBJECT_CODES:
            continue
        code = SUBJECT_CODES[name.strip()]
        for i, col in enumerate(ASSESSMENT_COLS):
            title = cell(col, row)
            if not title:
                continue
            assessments.append({
                'subject': code, 'term': TERM, 'name': str(title).strip(), 'criterion': criterion(str(title)),
                'weight': cell(col, row + 1) or 0, 'grade': cell(col, row + 2), 'sortOrder': i,
            })
        for i, col in enumerate(EXAM_COLS):
            title = cell(col, row)
            if title and cell(col, row + 1) is not None:
                components.append({'subject': code, 'name': str(title).strip().replace('Essessment', 'Assessment'),
                                   'weight': cell(col, row + 1), 'sortOrder': i})

    for s in subjects:
        del s['row']
    out = {'version': 1, 'subjects': subjects, 'scenarios': scenarios,
           'assessments': assessments, 'examComponents': components}
    with open(dst, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f'{len(scenarios)} scénářů, {len(assessments)} assessmentů, {len(components)} složek zkoušek → {dst}')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
